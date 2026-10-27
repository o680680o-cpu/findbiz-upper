#!/usr/bin/env python3
"""查指定公司與其上層關係公司、人員。同一筆資料只留一列。"""

from __future__ import annotations

import csv
import io
import os
import re
import threading
import time
from pathlib import Path
from typing import Any
from urllib.parse import quote

import requests
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

G0V = "https://company.g0v.ronny.tw/api"
GCIS = "https://data.gcis.nat.gov.tw/od/data/api"
GCIS_COMPANY = "5F64D864-61CB-4D0D-8AD9-492047CC1EA6"
GCIS_DIRECTORS = "4E5F7653-1B91-4DDC-99D5-468530FAE396"
FINDBIZ_DETAIL = "https://findbiz.nat.gov.tw/fts/query/QueryCmpyDetail/queryCmpyDetail.do"
TRADE_URL = "https://fbfh.trade.gov.tw/opendata/companyData.csv"
TRADE_INDEX = Path(__file__).resolve().parent / "data" / "trade_english.tsv"
TRADE_META = TRADE_INDEX.with_suffix(".meta")
MAX_LEVEL = 8
MAX_COMPANIES = 80
MAX_BATCH = 15
REQUEST_GAP = 0.2
IDENTITY_HEADERS = ("層級", "統一編號", "公司名稱", "登記現況", "來自查詢")

app = Flask(__name__)
_cache: dict[str, Any] = {}
_cache_lock = threading.Lock()
_throttle_lock = threading.Lock()
_trade_lock = threading.Lock()
_trade_names: dict[str, str] | None = None
_trade_refresh_started = False
_last_request = 0.0
_session = requests.Session()
_session.headers["User-Agent"] = "findbiz-upper/1.0"


def _throttle() -> None:
    global _last_request
    with _throttle_lock:
        wait = REQUEST_GAP - (time.time() - _last_request)
        if wait > 0:
            time.sleep(wait)
        _last_request = time.time()


def _get_json(url: str, params: dict | None = None) -> Any:
    _throttle()
    response = _session.get(url, params=params, timeout=30)
    response.raise_for_status()
    return response.json()


def normalize_id(value: Any) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) != 8 or set(digits) == {"0"}:
        return ""
    return digits


def text_parts(value: Any) -> list[str]:
    if isinstance(value, str):
        text = value.strip()
        return [text] if text else []
    if isinstance(value, list):
        return [item.strip() for item in value if isinstance(item, str) and item.strip()]
    return []


def looks_latin(text: str) -> bool:
    letters = [char for char in text if char.isalpha()]
    if not letters:
        return False
    latin = sum(char.isascii() for char in letters)
    return latin / len(letters) >= 0.6


def display_name(item: dict) -> str:
    for key in ("公司名稱", "商業名稱", "名稱"):
        parts = text_parts(item.get(key))
        if parts:
            return parts[0]
    finance = item.get("財政部") or {}
    if isinstance(finance, dict):
        parts = text_parts(finance.get("營業人名稱"))
        if parts:
            return parts[0]
    return ""


def english_name_of(item: dict) -> str:
    """章程所訂外文公司名稱。沒有登記就回傳空字串。"""
    for part in text_parts(item.get("章程所訂外文公司名稱")):
        return part
    parts = text_parts(item.get("公司名稱"))
    if len(parts) >= 2 and looks_latin(parts[1]) and parts[1] != parts[0]:
        return parts[1]
    return ""


def _read_trade_index() -> dict[str, str]:
    names: dict[str, str] = {}
    if not TRADE_INDEX.is_file():
        return names
    with TRADE_INDEX.open(encoding="utf-8") as handle:
        for line in handle:
            company_id, _, english = line.rstrip("\n").partition("\t")
            company_id = normalize_id(company_id)
            english = english.strip()
            if company_id and english:
                names[company_id] = english
    return names


def _build_trade_index() -> None:
    """把出進口廠商登記整理成統編對英文名稱。完整檔約 100MB，只留這兩欄。"""
    TRADE_INDEX.parent.mkdir(parents=True, exist_ok=True)
    raw_path = TRADE_INDEX.with_suffix(".csv.part")
    tmp_path = TRADE_INDEX.with_suffix(".tsv.part")
    try:
        with requests.get(
            TRADE_URL,
            timeout=180,
            stream=True,
            headers={"User-Agent": "findbiz-upper/1.0"},
        ) as response:
            response.raise_for_status()
            modified = (response.headers.get("Last-Modified") or "").strip()
            with raw_path.open("wb") as handle:
                for chunk in response.iter_content(256 * 1024):
                    if chunk:
                        handle.write(chunk)
        with raw_path.open(encoding="utf-8-sig", newline="") as source, tmp_path.open(
            "w", encoding="utf-8"
        ) as dest:
            for row in csv.DictReader(source):
                company_id = normalize_id(row.get("統一編號"))
                english = (row.get("廠商英文名稱") or "").strip()
                chinese = (row.get("廠商中文名稱") or "").strip()
                if not company_id or not english or not looks_latin(english):
                    continue
                if english.casefold() == chinese.casefold():
                    continue
                english = english.replace("\t", " ").replace("\r", " ").replace("\n", " ")
                dest.write(f"{company_id}\t{english}\n")
        tmp_path.replace(TRADE_INDEX)
        TRADE_META.write_text(modified + "\n", encoding="utf-8")
    finally:
        raw_path.unlink(missing_ok=True)
        tmp_path.unlink(missing_ok=True)


def _refresh_trade_index() -> None:
    global _trade_names
    if os.environ.get("TRADE_REFRESH") == "0":
        return
    try:
        if not TRADE_INDEX.is_file():
            _build_trade_index()
        else:
            response = requests.head(
                TRADE_URL,
                timeout=20,
                allow_redirects=True,
                headers={"User-Agent": "findbiz-upper/1.0"},
            )
            response.raise_for_status()
            remote = (response.headers.get("Last-Modified") or "").strip()
            local = TRADE_META.read_text(encoding="utf-8").strip() if TRADE_META.is_file() else ""
            if not remote or remote == local:
                return
            _build_trade_index()
        names = _read_trade_index()
        with _trade_lock:
            _trade_names = names
    except (OSError, requests.RequestException, csv.Error):
        return


def trade_english_name(company_id: Any) -> str:
    global _trade_names
    company_id = normalize_id(company_id)
    if not company_id:
        return ""
    names = _trade_names
    if names is None:
        with _trade_lock:
            if _trade_names is None:
                _trade_names = _read_trade_index()
            names = _trade_names
    return names.get(company_id, "")


def english_name_for(item: dict) -> str:
    """章程外文名稱優先。章程空白時，用出進口廠商英文名稱。"""
    charter = english_name_of(item)
    if charter:
        return charter
    return trade_english_name(item.get("統一編號"))


def entity_kind(item: dict) -> str:
    if item.get("公司名稱"):
        return "公司"
    if item.get("商業名稱"):
        return "商業"
    return "其他"


def status_of(item: dict) -> str:
    for key in ("登記現況", "公司狀況", "現況"):
        value = item.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def representative_of(item: dict) -> str:
    for key in ("代表人姓名", "負責人姓名"):
        value = item.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def location_of(item: dict) -> str:
    for key in ("公司所在地", "地址"):
        value = item.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def _masked_name(name: str) -> bool:
    return any(mark in name for mark in ("○", "◯", "〇", "●", "*", "＊"))


def _gcis_rows(dataset: str, company_id: str) -> list[dict]:
    payload = _get_json(
        f"{GCIS}/{dataset}",
        {
            "$format": "json",
            "$filter": f"Business_Accounting_NO eq {company_id}",
            "$skip": "0",
            "$top": "200",
        },
    )
    if not isinstance(payload, list):
        return []
    return [row for row in payload if isinstance(row, dict)]


def _apply_official_officers(data: dict, company_id: str) -> None:
    """g0v 的董監事會落後。經濟部這份每小時更新，有資料就改用它。"""
    try:
        profile = _gcis_rows(GCIS_COMPANY, company_id)
    except requests.RequestException:
        profile = []
    if profile:
        representative = str(profile[0].get("Responsible_Name") or "").strip()
        if representative and not _masked_name(representative):
            data["代表人姓名"] = representative
        status = str(profile[0].get("Company_Status_Desc") or "").strip()
        if status:
            data["登記現況"] = status
    try:
        directors = _gcis_rows(GCIS_DIRECTORS, company_id)
    except requests.RequestException:
        return
    rows = []
    for item in directors:
        person = str(item.get("Person_Name") or "").strip()
        if not person or _masked_name(person):
            return
        rows.append({
            "姓名": person,
            "職稱": str(item.get("Person_Position_Name") or "董監事").strip() or "董監事",
            "出資額": item.get("Person_Shareholding"),
            "所代表法人": str(item.get("Juristic_Person_Name") or "").strip(),
        })
    if rows:
        data["董監事名單"] = rows


def show_company(company_id: str) -> dict | None:
    company_id = normalize_id(company_id)
    if not company_id:
        return None
    with _cache_lock:
        cached = _cache.get(company_id)
    if cached is not None:
        return cached or None
    try:
        payload = _get_json(f"{G0V}/show/{company_id}")
        data = payload.get("data") if isinstance(payload, dict) else None
    except requests.RequestException:
        data = None
    if isinstance(data, dict):
        data["統一編號"] = company_id
        _apply_official_officers(data, company_id)
    with _cache_lock:
        _cache[company_id] = data or {}
    return data


def resolve_names(names: list[str]) -> dict[str, str]:
    cleaned = []
    seen = set()
    for name in names:
        name = (name or "").strip()
        if name and name not in seen:
            seen.add(name)
            cleaned.append(name)
    found: dict[str, str] = {}
    for start in range(0, len(cleaned), 20):
        chunk = cleaned[start:start + 20]
        try:
            rows = _get_json(f"{G0V}/bulksearch", {"names": ",".join(chunk)})
        except requests.RequestException:
            continue
        if not isinstance(rows, list):
            continue
        for row in rows:
            query = (row.get("query") or "").strip()
            company_id = normalize_id(row.get("result"))
            if query and company_id:
                found[query] = company_id
    return found


def search_companies(keyword: str, limit: int = 10) -> dict:
    keyword = keyword.strip()
    try:
        payload = _get_json(f"{G0V}/search", {"q": keyword, "page": 1})
    except requests.RequestException as exc:
        raise LookupError("公司資料服務暫時連不上，請稍後再試") from exc
    items = payload.get("data") if isinstance(payload, dict) else []
    found = int(payload.get("found") or len(items or []))
    results = []
    seen = set()
    for item in items or []:
        if not isinstance(item, dict):
            continue
        company_id = normalize_id(item.get("統一編號") or item.get("公司統編"))
        if not company_id or company_id in seen:
            continue
        seen.add(company_id)
        name = display_name(item)
        if not name:
            full = show_company(company_id)
            if full:
                item = {**item, **full}
                name = display_name(item)
        results.append({
            "company_id": company_id,
            "company_name": name or f"統編 {company_id}",
            "status": status_of(item),
            "location": location_of(item),
            "representative": representative_of(item),
            "entity_type": entity_kind(item),
        })

    folded = keyword.casefold()

    def rank(entry: dict) -> tuple:
        name = entry["company_name"].casefold()
        if name == folded:
            closeness = 0
        elif name.startswith(folded):
            closeness = 1
        elif folded in name:
            closeness = 2
        else:
            closeness = 3
        kind = 0 if entry["entity_type"] == "公司" else 1
        closed = 1 if any(word in entry["status"] for word in ("解散", "撤銷", "廢止", "停業")) else 0
        return (closeness, kind, closed, name)

    results.sort(key=rank)
    return {"keyword": keyword, "found": found, "results": results[:limit]}


def _juristic(value: Any) -> tuple[str, str]:
    if isinstance(value, list) and len(value) >= 2:
        return normalize_id(value[0]), str(value[1]).strip()
    if isinstance(value, str):
        return "", value.strip()
    return "", ""


def _parents_of(data: dict) -> list[dict]:
    grouped: dict[str, dict] = {}
    for director in data.get("董監事名單") or []:
        corp_id, corp_name = _juristic(director.get("所代表法人"))
        if not corp_name:
            continue
        key = corp_id or corp_name
        slot = grouped.setdefault(key, {
            "id": corp_id,
            "name": corp_name,
            "people": [],
        })
        person = (director.get("姓名") or "").strip()
        role = (director.get("職稱") or "董監事").strip()
        if person and (person, role) not in slot["people"]:
            slot["people"].append((person, role))
    pending = [item["name"] for item in grouped.values() if not item["id"]]
    resolved = resolve_names(pending) if pending else {}
    parents = []
    seen_ids = set()
    for item in grouped.values():
        corp_id = item["id"] or resolved.get(item["name"], "")
        if corp_id and corp_id in seen_ids:
            continue
        if corp_id:
            seen_ids.add(corp_id)
        item["id"] = corp_id
        parents.append(item)
    return parents


def _people_of(data: dict, company_id: str, company_name: str, level: int) -> list[dict]:
    people = []
    seen = set()

    def add(name: str, role: str, shareholding: str = "", juristic: str = "") -> None:
        name = (name or "").strip()
        role = (role or "").strip()
        if not name:
            return
        key = (name, role, juristic)
        if key in seen:
            return
        seen.add(key)
        people.append({
            "name": name,
            "role": role,
            "shareholding": shareholding,
            "juristic": juristic,
            "company_id": company_id,
            "company_name": company_name,
            "level": level,
        })

    for director in data.get("董監事名單") or []:
        _, corp_name = _juristic(director.get("所代表法人"))
        share = director.get("出資額")
        add(
            director.get("姓名", ""),
            director.get("職稱", "董監事"),
            "" if share is None else str(share),
            corp_name,
        )
    for manager in data.get("經理人名單") or []:
        add(manager.get("姓名", ""), manager.get("職稱") or "經理人")
    add(representative_of(data), "負責人")
    return people


def _parse_depth(raw: str) -> int:
    if raw == "all":
        return MAX_LEVEL
    try:
        return max(0, min(MAX_LEVEL, int(raw)))
    except ValueError:
        return MAX_LEVEL


def query_upper(company_id: str, depth: str) -> dict:
    company_id = normalize_id(company_id)
    if not company_id:
        raise ValueError("請輸入 8 碼統一編號")
    max_level = _parse_depth(depth)
    root = show_company(company_id)
    if not root:
        raise LookupError(f"找不到統編 {company_id} 的公司資料")

    companies: dict[str, dict] = {}
    loose: dict[str, dict] = {}
    appearances: list[dict] = []
    notes: list[str] = []
    truncated = False

    def ensure_company(cid: str, data: dict, level: int, via: str) -> dict:
        row = companies.get(cid)
        if row is None:
            row = {
                "level": level,
                "company_id": cid,
                "company_name": display_name(data) or f"統編 {cid}",
                "status": status_of(data),
                "representative": representative_of(data),
                "location": location_of(data),
                "english_name": english_name_for(data),
                "via": [],
                "findbiz_url": FINDBIZ_DETAIL,
            }
            companies[cid] = row
        elif not row.get("english_name"):
            row["english_name"] = english_name_for(data)
        if level < row["level"]:
            row["level"] = level
        if via and via not in row["via"]:
            row["via"].append(via)
        return row

    ensure_company(company_id, root, 0, "查詢公司")
    appearances.extend(_people_of(root, company_id, display_name(root), 0))

    queue = [(company_id, 0)]
    queued = {company_id}
    while queue:
        current_id, level = queue.pop(0)
        if level >= max_level:
            continue
        current = show_company(current_id)
        if not current:
            continue
        child_name = display_name(current) or current_id
        for parent in _parents_of(current):
            parent_id = parent["id"]
            parent_name = parent["name"]
            spokes = "、".join(f"{person}（{role}）" for person, role in parent["people"][:3])
            extra = len(parent["people"]) - 3
            if extra > 0:
                spokes += f" 等 {len(parent['people'])} 人"
            via = f"由 {spokes} 代表，出資給 {child_name}" if spokes else f"出資給 {child_name}"
            if not parent_id:
                slot = loose.setdefault(parent_name, {
                    "level": level + 1,
                    "company_id": "",
                    "company_name": parent_name,
                    "status": "查無公司統編",
                    "representative": "",
                    "location": "",
                    "english_name": "",
                    "via": [],
                    "findbiz_url": "",
                })
                if via not in slot["via"]:
                    slot["via"].append(via)
                if level + 1 < slot["level"]:
                    slot["level"] = level + 1
                continue
            if parent_id == current_id:
                continue
            if len(companies) >= MAX_COMPANIES and parent_id not in companies:
                if not truncated:
                    truncated = True
                    notes.append(f"公司數量已達 {MAX_COMPANIES} 家，其餘上層未再展開")
                continue
            pdata = show_company(parent_id)
            if not pdata:
                notes.append(f"查無 {parent_name}（{parent_id}）的登記明細，仍列入名單")
                pdata = {"公司名稱": parent_name, "統一編號": parent_id}
            ensure_company(parent_id, pdata, level + 1, via)
            appearances.extend(_people_of(
                pdata, parent_id, display_name(pdata) or parent_name, level + 1,
            ))
            if parent_id not in queued and level + 1 < max_level:
                queued.add(parent_id)
                queue.append((parent_id, level + 1))

    if depth == "all" and any(row["level"] >= MAX_LEVEL for row in companies.values()):
        notes.append(f"往上已追溯到第 {MAX_LEVEL} 層")

    company_rows = list(companies.values()) + list(loose.values())
    company_rows.sort(key=lambda row: (row["level"], row["company_name"]))
    for row in company_rows:
        row["via_text"] = "；".join(row.pop("via"))
        row["key"] = row["company_id"] or f"name:{row['company_name']}"

    assignments = [
        {
            "key": item["company_id"] or f"name:{item['company_name']}",
            "name": item["name"],
            "role": item["role"],
        }
        for item in appearances
        if item.get("name")
    ]
    return {
        "root_id": company_id,
        "root_name": display_name(root),
        "depth": depth,
        "truncated": truncated,
        "notes": notes,
        "companies": company_rows,
        "assignments": assignments,
    }


def _unique_ids(raw_ids: list) -> list[str]:
    ids: list[str] = []
    seen = set()
    for value in raw_ids:
        company_id = normalize_id(value)
        if not company_id or company_id in seen:
            continue
        seen.add(company_id)
        ids.append(company_id)
    return ids


def query_many(company_ids: list[str], depth: str) -> dict:
    """查多家公司，公司與人員都去重，組成「列是公司、欄是人名」的對照表。"""
    if not company_ids:
        raise ValueError("請至少輸入一家公司")
    if len(company_ids) > MAX_BATCH:
        raise ValueError(f"一次最多查 {MAX_BATCH} 家公司")

    merged: dict[str, dict] = {}
    ordered_names: dict[str, list[str]] = {}
    notes: list[str] = []
    roots: list[dict[str, str]] = []
    truncated = False

    for company_id in company_ids:
        one = query_upper(company_id, depth)
        root_name = one["root_name"]
        roots.append({"id": one["root_id"], "name": root_name})
        if one["truncated"]:
            truncated = True
        for note in one["notes"]:
            if note not in notes:
                notes.append(note)
        for company in one["companies"]:
            key = company["key"]
            current = merged.get(key)
            if current is None:
                current = dict(company)
                current["sources"] = [root_name]
                merged[key] = current
            else:
                if root_name not in current["sources"]:
                    current["sources"].append(root_name)
                if not current.get("english_name") and company.get("english_name"):
                    current["english_name"] = company["english_name"]
                if company["level"] < current["level"]:
                    current["level"] = company["level"]
                    current["via_text"] = company["via_text"]
                elif (
                    company["via_text"]
                    and company["via_text"] not in (current.get("via_text") or "")
                    and current.get("via_text") not in ("", "查詢公司")
                ):
                    current["via_text"] += "；" + company["via_text"]
        for item in one["assignments"]:
            if item.get("role") == "負責人":
                continue
            _append_name(ordered_names, item["key"], item["name"])

    grid = []
    widest = 1
    for company in sorted(merged.values(), key=lambda row: (row["level"], row["company_name"])):
        names = _names_for_company(company, ordered_names.get(company["key"], []))
        widest = max(widest, len(names))
        grid.append({
            "level": company["level"],
            "company_id": company["company_id"],
            "company_name": company["company_name"],
            "status": company["status"],
            "representative": company["representative"],
            "via_text": company["via_text"],
            "sources": company["sources"],
            "english_name": company.get("english_name") or "",
            "names": names,
        })

    people = ["負責人"] + [str(index) for index in range(2, widest + 1)]
    for record in grid:
        record["cells"] = record["names"] + [""] * (widest - len(record["names"]))

    levels = max((row["level"] for row in grid), default=0)
    return {
        "roots": roots,
        "depth": depth,
        "truncated": truncated,
        "notes": notes,
        "people": people,
        "grid": grid,
        "stats": {
            "queries": len(roots),
            "companies": len(grid),
            "people": len(people),
            "levels": levels,
        },
    }


def _append_name(ordered: dict[str, list[str]], key: str, name: str) -> None:
    name = (name or "").strip()
    if not name:
        return
    names = ordered.setdefault(key, [])
    if name not in names:
        names.append(name)


def _names_for_company(company: dict, others: list[str]) -> list[str]:
    """第一格是負責人，後面依登記順序接其餘姓名，同一人不重複。"""
    representative = (company.get("representative") or "").strip()
    rest = [name for name in others if name and name != representative]
    return [representative, *rest]


def _level_label(level: int) -> str:
    if level == 0:
        return "查詢公司"
    return f"上層第 {level} 層"


def build_workbook(result: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "公司與人員"
    headers = list(IDENTITY_HEADERS) + list(result["people"])
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="8F2D1F")
    header_font = Font(color="FFFFFF", bold=True, name="PingFang TC", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="E4DCD0"),
        right=Side(style="thin", color="E4DCD0"),
        top=Side(style="thin", color="E4DCD0"),
        bottom=Side(style="thin", color="E4DCD0"),
    )
    query_fill = PatternFill("solid", fgColor="FBF6EF")

    for column, _header in enumerate(headers, 1):
        cell = sheet.cell(1, column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    for record in result["grid"]:
        sources = "、".join(record.get("sources") or [])
        sheet.append([
            _level_label(record["level"]),
            record["company_id"],
            record["company_name"],
            record["status"],
            sources,
            *record["cells"],
        ])
        row_number = sheet.max_row
        for column in range(1, len(headers) + 1):
            cell = sheet.cell(row_number, column)
            cell.border = thin
            cell.alignment = left if column in (3, 5) else center
            if record["level"] == 0 and column <= len(IDENTITY_HEADERS):
                cell.fill = query_fill
            if column > len(IDENTITY_HEADERS) and cell.value:
                cell.font = Font(name="PingFang TC", size=11)
        english = (record.get("english_name") or "").strip()
        if english and english.casefold() != (record.get("company_name") or "").casefold():
            sheet.append([""] * len(headers))
            sheet.cell(sheet.max_row, 3, english)
            english_row = sheet.max_row
            for column in range(1, len(headers) + 1):
                cell = sheet.cell(english_row, column)
                cell.border = thin
                cell.font = Font(italic=True, color="6B6258", name="Calibri", size=11)
                cell.alignment = left if column == 3 else center

    sheet.freeze_panes = "G2"
    sheet.row_dimensions[1].height = 32
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    widths = (14, 14, 52, 14, 28)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for column in range(len(IDENTITY_HEADERS) + 1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 12
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddHeader.left.text = "列是公司。第一個姓名是負責人，後面依序是其餘人員，不列職稱。"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _export_filename(result: dict) -> str:
    roots = result["roots"]
    if len(roots) == 1:
        stem = roots[0]["name"] or roots[0]["id"]
    else:
        stem = f"{roots[0]['name']}等{len(roots)}家"
    stem = re.sub(r'[\\/:*?"<>|]', "", stem).strip() or "公司人員對照"
    return f"{stem}_公司人員對照.xlsx"


def _read_query_body() -> tuple[list[str], str]:
    if request.method == "POST":
        body = request.get_json(silent=True) or {}
        raw_ids = body.get("ids") or []
        depth = str(body.get("depth") or "all")
    else:
        raw = request.args.get("ids") or request.args.get("id") or ""
        raw_ids = re.split(r"[\s,，、;；]+", raw)
        depth = request.args.get("depth", "all")
    if isinstance(raw_ids, str):
        raw_ids = re.split(r"[\s,，、;；]+", raw_ids)
    return _unique_ids(list(raw_ids)), depth


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/search")
def api_search():
    keyword = request.args.get("q", "").strip()
    if not keyword:
        return jsonify({"error": "請輸入公司名稱或統一編號"}), 400
    if re.fullmatch(r"\d{8}", keyword):
        data = show_company(keyword)
        if not data:
            return jsonify({"error": f"找不到統編 {keyword}", "results": []}), 404
        return jsonify({
            "keyword": keyword,
            "found": 1,
            "results": [{
                "company_id": keyword,
                "company_name": display_name(data) or f"統編 {keyword}",
                "status": status_of(data),
                "location": location_of(data),
                "representative": representative_of(data),
                "entity_type": "公司",
            }],
        })
    try:
        return jsonify(search_companies(keyword))
    except LookupError as exc:
        return jsonify({"error": str(exc), "results": []}), 502


@app.route("/api/query", methods=["GET", "POST"])
def api_query():
    company_ids, depth = _read_query_body()
    try:
        return jsonify(query_many(company_ids, depth))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except LookupError as exc:
        return jsonify({"error": str(exc)}), 404


@app.route("/api/export", methods=["POST"])
def api_export():
    company_ids, depth = _read_query_body()
    try:
        result = query_many(company_ids, depth)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except LookupError as exc:
        return jsonify({"error": str(exc)}), 404

    filename = _export_filename(result)
    response = send_file(
        io.BytesIO(build_workbook(result)),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )
    encoded = quote(filename)
    response.headers["Content-Disposition"] = (
        f"attachment; filename=\"companies.xlsx\"; filename*=UTF-8''{encoded}"
    )
    return response


@app.route("/health")
def health():
    return jsonify({"ok": True})


def _boot_trade_index() -> None:
    global _trade_names, _trade_refresh_started
    with _trade_lock:
        if _trade_names is None:
            _trade_names = _read_trade_index()
        if _trade_refresh_started:
            return
        _trade_refresh_started = True
    if os.environ.get("TRADE_REFRESH") == "0":
        return
    threading.Thread(target=_refresh_trade_index, name="trade-english", daemon=True).start()


_boot_trade_index()


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5070, debug=False)
